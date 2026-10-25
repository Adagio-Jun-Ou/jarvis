"""
Office-document generation tools.

The agent can produce real, downloadable Microsoft Office artifacts —
Excel workbooks (.xlsx), PowerPoint decks (.pptx), and Word documents
(.docx) — plus read spreadsheets back.  Every file is written into the
user's ``outputs/`` folder (via :func:`agent.workspace_tools.output_path`)
so it appears in the Files tab and is downloadable through
``GET /api/files/download``.

The functions follow the tool-registry calling convention: the VM
injects ``user_email``/``instance_scope``; the LLM supplies the content.
They accept forgiving, JSON-friendly argument shapes because small
models are imprecise, and they degrade with a clear error (never a
raw traceback) when a library or input is missing.
"""

from __future__ import annotations

import logging
from typing import Any

log = logging.getLogger("agent.tools.office")


def _slug(name: str, ext: str) -> str:
    """Sanitize a filename to ``<safe-stem><ext>``.  Preserves an
    existing correct extension and strips any other one the model may
    have appended, so ``tax_note.docx`` and ``tax_note`` both become
    ``tax_note.docx`` (not ``tax_notedocx.docx``)."""
    raw = (name or "").strip().replace("\\", "/").rsplit("/", 1)[-1]
    stem = raw[: -len(ext)] if raw.lower().endswith(ext) else raw
    # Drop any other trailing extension the model tacked on.
    if "." in stem:
        head, tail = stem.rsplit(".", 1)
        if tail.lower() in ("xlsx", "pptx", "docx", "csv", "pdf", "xls", "ppt", "doc"):
            stem = head
    safe = "".join(c for c in stem if c.isalnum() or c in " -_").strip()
    safe = (safe or "document").replace(" ", "_")[:80]
    return f"{safe}{ext}"


def _download(user_email: str, filename: str, writer) -> dict:
    """Resolve the output path, run ``writer(path)``, return a result
    dict with a download URL the UI/agent can hand to the user."""
    from agent.workspace_tools import output_path

    try:
        path = output_path(user_email, filename)
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": f"could not resolve output path: {exc}"}
    try:
        writer(path)
    except Exception as exc:  # noqa: BLE001
        log.exception("office write failed")
        return {"ok": False, "error": str(exc)}
    rel = f"outputs/{filename}"
    return {
        "ok": True,
        "path": rel,
        "download": f"/api/files/download?path={rel}",
        "message": f"Created {filename}. It's in your Files and ready to download.",
    }


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------


def create_spreadsheet(
    filename: str,
    sheets: Any = None,
    headers: Any = None,
    rows: Any = None,
    *,
    user_email: str,
    instance_scope: str = "default",
) -> dict:
    """Create an Excel (.xlsx) workbook. Single tab: pass ``headers`` (a
    list of column names) and ``rows`` (a list of rows, where EACH row is
    itself a list of cell values), e.g. headers=["Month","Revenue"],
    rows=[["Jan",100],["Feb",150]]. Multiple tabs: pass ``sheets`` as a
    list of {name, headers, rows}. Numbers are stored as numbers; the
    header row is bold. Row shapes are auto-normalized, but a list of
    per-row lists is best."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils import get_column_letter
    except Exception:  # noqa: BLE001
        return {"ok": False, "error": "openpyxl is not installed on this VM"}

    if not sheets:
        sheets = [{"name": "Sheet1", "headers": headers or [], "rows": rows or []}]
    if isinstance(sheets, dict):
        sheets = [sheets]

    def _num(v):
        if isinstance(v, (int, float)):
            return v
        s = str(v).strip().replace(",", "").replace("$", "")
        try:
            return int(s) if s.isdigit() or (s.startswith("-") and s[1:].isdigit()) else float(s)
        except (ValueError, AttributeError):
            # Formula injection: openpyxl types any string beginning
            # with "=" as a formula, so text scraped from an email or
            # web page would become a live formula (=WEBSERVICE...)
            # the moment the user opened the workbook. Text stays text.
            if isinstance(v, str) and v.lstrip().startswith("="):
                return "'" + v
            return v

    def _norm_list(x):
        """Best-effort turn ``x`` into a python list. Accepts a real
        list, a JSON string, or a comma/newline-delimited string."""
        if x is None:
            return []
        if isinstance(x, (list, tuple)):
            return list(x)
        if isinstance(x, str):
            s = x.strip()
            if not s:
                return []
            if s[0] in "[{":
                try:
                    import json as _json

                    v = _json.loads(s)
                    return v if isinstance(v, list) else [v]
                except Exception:  # noqa: BLE001
                    pass
            # CSV-ish: newlines are rows, commas are cells.
            if "\n" in s:
                return [line for line in s.splitlines() if line.strip()]
            return [p.strip() for p in s.split(",")]
        return [x]

    def _rows_matrix(hdrs, body):
        """Normalize ``body`` into a list of row-lists. Handles a JSON/CSV
        string, a list of dicts (keyed by headers), a proper list of
        rows, and a flat list (chunked by header count)."""
        body = _norm_list(body)
        ncol = len(hdrs) if hdrs else 0
        out = []
        # list of dicts -> order by headers (or the dict's own keys)
        if body and all(isinstance(r, dict) for r in body):
            keys = hdrs or list(body[0].keys())
            return [[r.get(k, "") for k in keys] for r in body]
        # list of rows (each a list) OR CSV row-strings
        rowish = []
        for r in body:
            if isinstance(r, (list, tuple)):
                rowish.append(list(r))
            elif isinstance(r, str) and ("," in r):
                rowish.append([c.strip() for c in r.split(",")])
            else:
                rowish.append(r)
        # If everything is a scalar and we know the column count, chunk.
        if ncol > 1 and all(not isinstance(r, list) for r in rowish):
            return [rowish[i:i + ncol] for i in range(0, len(rowish), ncol)]
        for r in rowish:
            out.append(r if isinstance(r, list) else [r])
        return out

    def write(path: str) -> None:
        wb = Workbook()
        wb.remove(wb.active)
        for i, sh in enumerate(sheets):
            sh = sh if isinstance(sh, dict) else {}
            ws = wb.create_sheet(str(sh.get("name") or f"Sheet{i + 1}")[:31])
            hdrs = _norm_list(sh.get("headers"))
            body = _rows_matrix(hdrs, sh.get("rows"))
            r = 1
            if hdrs:
                for c, h in enumerate(hdrs, 1):
                    cell = ws.cell(row=1, column=c, value=str(h))
                    cell.font = Font(bold=True)
                r = 2
            widths: dict[int, int] = {}
            for row in body:
                row = row if isinstance(row, (list, tuple)) else [row]
                for c, v in enumerate(row, 1):
                    ws.cell(row=r, column=c, value=_num(v))
                    widths[c] = max(widths.get(c, 8), min(48, len(str(v)) + 2))
                r += 1
            for c, w in widths.items():
                ws.column_dimensions[get_column_letter(c)].width = w
        wb.save(path)

    return _download(user_email, _slug(filename, ".xlsx"), write)


def read_spreadsheet(
    path: str,
    max_rows: int = 200,
    *,
    user_email: str,
    instance_scope: str = "default",
) -> dict:
    """Read a spreadsheet from the user's workspace and return its rows
    (first sheet). Use this before summarizing or transforming data."""
    try:
        from openpyxl import load_workbook
    except Exception:  # noqa: BLE001
        return {"ok": False, "error": "openpyxl is not installed on this VM"}
    from agent.workspace_tools import _resolve_under, _user_root

    try:
        p = _resolve_under(_user_root(user_email), path)
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": str(exc)}
    try:
        wb = load_workbook(p, read_only=True, data_only=True)
        ws = wb.active
        out = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max(1, min(int(max_rows), 2000)):
                break
            out.append([("" if v is None else v) for v in row])
        return {"ok": True, "sheet": ws.title, "rows": out, "truncated": ws.max_row > len(out)}
    except Exception as exc:  # noqa: BLE001
        return {"ok": False, "error": str(exc)}


# ---------------------------------------------------------------------------
# PowerPoint
# ---------------------------------------------------------------------------


def create_presentation(
    filename: str,
    title: str = "",
    slides: Any = None,
    *,
    user_email: str,
    instance_scope: str = "default",
) -> dict:
    """Create a PowerPoint (.pptx) deck. ``slides`` is a list of
    {title, bullets: [...]} (or {title, body}). A title slide is added
    from ``title`` when provided."""
    try:
        from pptx import Presentation
        from pptx.util import Pt
    except Exception:  # noqa: BLE001
        return {"ok": False, "error": "python-pptx is not installed on this VM"}

    slides = slides or []
    if isinstance(slides, dict):
        slides = [slides]

    def write(path: str) -> None:
        prs = Presentation()
        if title:
            s = prs.slides.add_slide(prs.slide_layouts[0])
            s.shapes.title.text = str(title)
            if len(s.placeholders) > 1:
                s.placeholders[1].text = "Prepared by Alex"
        for sl in slides:
            sl = sl if isinstance(sl, dict) else {"title": str(sl)}
            layout = prs.slide_layouts[1]
            s = prs.slides.add_slide(layout)
            s.shapes.title.text = str(sl.get("title") or "")
            bullets = sl.get("bullets")
            if bullets is None and sl.get("body"):
                bullets = [ln for ln in str(sl["body"]).split("\n") if ln.strip()]
            body = s.placeholders[1].text_frame if len(s.placeholders) > 1 else None
            if body is not None and bullets:
                body.clear()
                for j, b in enumerate(bullets):
                    para = body.paragraphs[0] if j == 0 else body.add_paragraph()
                    para.text = str(b)
                    para.font.size = Pt(18)
        prs.save(path)

    return _download(user_email, _slug(filename, ".pptx"), write)


# ---------------------------------------------------------------------------
# Word
# ---------------------------------------------------------------------------


def create_document(
    filename: str,
    title: str = "",
    body: Any = None,
    sections: Any = None,
    *,
    user_email: str,
    instance_scope: str = "default",
) -> dict:
    """Create a Word (.docx) document. Provide ``body`` (text, blank
    lines separate paragraphs) and/or ``sections`` (list of
    {heading, text})."""
    try:
        from docx import Document
    except Exception:  # noqa: BLE001
        return {"ok": False, "error": "python-docx is not installed on this VM"}

    sections = sections or []
    if isinstance(sections, dict):
        sections = [sections]

    def write(path: str) -> None:
        doc = Document()
        if title:
            doc.add_heading(str(title), level=0)
        if body:
            for para in str(body).split("\n\n"):
                if para.strip():
                    doc.add_paragraph(para.strip())
        for sec in sections:
            sec = sec if isinstance(sec, dict) else {"text": str(sec)}
            if sec.get("heading"):
                doc.add_heading(str(sec["heading"]), level=1)
            if sec.get("text"):
                for para in str(sec["text"]).split("\n\n"):
                    if para.strip():
                        doc.add_paragraph(para.strip())
        doc.save(path)

    return _download(user_email, _slug(filename, ".docx"), write)
